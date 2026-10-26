"""Helpers for the knowledge tests: a deterministic embedder, a fake Ollama HTTP server (the ingestion worker is a
separate process, so it needs a real socket), and generators for PDF, DOCX, CSV, Markdown and text files."""

from __future__ import annotations

import hashlib
import json
import math
import re
import threading
import time
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any, Callable

from xodus.retrieval.fts import STOPWORDS

DIM = 128


def hash_vector(text: str, dim: int = DIM) -> list[float]:
    """Bag-of-words hashing embedder: texts that share content words have similar vectors."""
    vec = [0.0] * dim
    for word in re.findall(r"\w+", text.lower()):
        if word in STOPWORDS:
            continue
        h = int.from_bytes(hashlib.blake2b(word.encode(), digest_size=8).digest(), "big")
        vec[h % dim] += 1.0 if (h >> 40) & 1 else -1.0
    norm = math.sqrt(sum(v * v for v in vec))
    return [v / norm for v in vec] if norm else [1.0] + [0.0] * (dim - 1)


class HashEmbedder:
    """Synchronous embedder for in-process pipeline tests; records every call."""

    def __init__(self, dim: int = DIM) -> None:
        self.dim = dim
        self.calls: list[list[str]] = []

    def embed(self, texts: list[str]) -> list[list[float]]:
        self.calls.append(list(texts))
        return [hash_vector(t, self.dim) for t in texts]

    @property
    def total_texts(self) -> int:
        return sum(len(c) for c in self.calls)


class FakeOllama:
    """A local HTTP server speaking enough of Ollama's API for XODUS: version, tags, show, embed and chat."""

    def __init__(self, dim: int = DIM) -> None:
        self.dim = dim
        self.embed_calls: list[dict[str, Any]] = []
        self.embed_delay_s = 0.0
        self.embed_fail = 0  # answer the next N embed calls with HTTP 503
        self.embed_hang_s = 0.0  # sleep this long before answering (to provoke a client timeout)
        self.embed_model_missing = False
        self.show_missing_for: set[str] = set()  # /api/show answers 404 for these model names
        self.chat_answer: Callable[[list[dict[str, Any]]], str] | None = None
        self.chat_requests: list[dict[str, Any]] = []
        self.on_embed: Callable[[list[str]], None] | None = None
        self._lock = threading.Lock()
        outer = self

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args: Any) -> None:
                pass

            def _json(self, status: int, body: Any) -> None:
                data = json.dumps(body).encode()
                self.send_response(status)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)

            def do_GET(self) -> None:  # noqa: N802
                if self.path == "/api/version":
                    self._json(200, {"version": "0.0-fake"})
                elif self.path == "/api/tags":
                    self._json(200, {"models": [{"name": "fake-model:latest", "digest": "abc", "size": 1},
                                                {"name": "fake-embed:latest", "digest": "def", "size": 1}]})
                else:
                    self._json(404, {"error": "not found"})

            def do_POST(self) -> None:  # noqa: N802
                length = int(self.headers.get("Content-Length", 0))
                body = json.loads(self.rfile.read(length) or b"{}")
                if self.path == "/api/show":
                    if body.get("model") in outer.show_missing_for:
                        return self._json(404, {"error": f"model '{body.get('model')}' not found"})
                    if body.get("model") == "fake-embed:latest" or "embed" in str(body.get("model")):
                        self._json(200, {"capabilities": ["embedding"], "details": {}, "model_info": {"fake.embedding_length": outer.dim}})
                    else:
                        self._json(200, {"capabilities": ["completion", "tools"], "details": {}, "model_info": {}})
                elif self.path == "/api/embed":
                    with outer._lock:
                        outer.embed_calls.append({"input": body.get("input"), "options": body.get("options"),
                                                  "dimensions": body.get("dimensions")})
                        fail = outer.embed_fail > 0
                        outer.embed_fail = max(0, outer.embed_fail - 1)
                    if outer.embed_model_missing:
                        return self._json(404, {"error": "model not found"})
                    if fail:
                        return self._json(503, {"error": "overloaded"})
                    if outer.embed_hang_s:
                        time.sleep(outer.embed_hang_s)
                    if outer.embed_delay_s:
                        time.sleep(outer.embed_delay_s)
                    texts = body.get("input") or []
                    if outer.on_embed:
                        outer.on_embed(texts)
                    dim = body.get("dimensions") or outer.dim
                    self._json(200, {"embeddings": [hash_vector(t, dim) for t in texts], "prompt_eval_count": 10 * len(texts)})
                elif self.path == "/api/chat":
                    with outer._lock:
                        outer.chat_requests.append(body)
                    text = outer.chat_answer(body["messages"]) if outer.chat_answer else "ok"
                    self.send_response(200)
                    self.send_header("Content-Type", "application/x-ndjson")
                    self.end_headers()
                    for word in text.split(" "):
                        self.wfile.write(json.dumps({"message": {"content": word + " "}, "done": False}).encode() + b"\n")
                    self.wfile.write(json.dumps({"done": True, "done_reason": "stop", "prompt_eval_count": 7, "eval_count": 3}).encode() + b"\n")
                else:
                    self._json(404, {"error": "not found"})

        self.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.server.daemon_threads = True
        self.port = self.server.server_address[1]
        self.url = f"http://127.0.0.1:{self.port}"
        self._thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self._thread.start()

    def stop(self) -> None:
        self.server.shutdown()
        self.server.server_close()

    def embedded_texts(self) -> list[str]:
        return [t for call in self.embed_calls for t in (call["input"] or [])]


# ---------------------------------------------------------------------------------------------- file generators


def pdf_bytes(pages: list[str]) -> bytes:
    """A small valid PDF with a text layer (Helvetica), one page per string; lines are separated by newlines."""

    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages)))
    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>",
                            f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode(),
                            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, text in enumerate(pages):
        stream = "BT /F1 11 Tf 14 TL 40 750 Td " + " ".join(f"({esc(line)}) Tj T*" for line in text.split("\n")) + " ET"
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {5 + 2 * i} 0 R "
                       f"/Resources << /Font << /F1 3 0 R >> >> >>".encode())
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream".encode())
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def write_pdf(path: Path, pages: list[str]) -> Path:
    path.write_bytes(pdf_bytes(pages))
    return path


def write_docx(path: Path, items: list[tuple[str, str]], table: list[list[str]] | None = None) -> Path:
    """items: (style, text) with style 'h1'..'h3' or 'p'. Built with python-docx (a real producer of DOCX files)."""
    import docx

    document = docx.Document()
    for style, text in items:
        if style.startswith("h"):
            document.add_heading(text, level=int(style[1:]))
        else:
            document.add_paragraph(text)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    document.save(str(path))
    return path


def write_csv(path: Path, header: list[str], rows: list[list[str]], delimiter: str = ",", encoding: str = "utf-8") -> Path:
    import csv
    import io

    buf = io.StringIO(newline="")
    writer = csv.writer(buf, delimiter=delimiter)
    writer.writerow(header)
    writer.writerows(rows)
    path.write_bytes(buf.getvalue().encode(encoding))
    return path


def zip_with(path: Path, entries: dict[str, bytes], compression: int = zipfile.ZIP_DEFLATED) -> Path:
    with zipfile.ZipFile(path, "w", compression) as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    return path


TOPICS = {
    "astronomy": "The telescope observed a distant nebula where newborn stars form inside enormous clouds of hydrogen gas. "
                 "Astronomers measured the redshift of the galaxy to estimate its distance and the expansion of the universe. "
                 "A supernova remnant glowed in X-ray light while the pulsar spun many times each second. ",
    "cooking": "Whisk the eggs with sugar until pale, then fold in the flour and melted butter for a tender sponge cake. "
               "Simmer the tomato sauce with garlic and basil, season with salt, and toss it through freshly boiled pasta. "
               "Rest the dough in the refrigerator overnight so the gluten relaxes before baking the sourdough loaf. ",
    "networking": "A router forwards packets between subnets using routing tables and the longest prefix match rule. "
                  "The TCP handshake establishes a connection with SYN and ACK segments before any application data flows. "
                  "DNS resolves hostnames to addresses while a firewall filters traffic by port and protocol. ",
    "gardening": "Plant tomato seedlings after the last frost and water the roots deeply in the early morning hours. "
                 "Mulch the vegetable beds with straw to suppress weeds, and prune the roses in late winter for healthy growth. "
                 "Compost kitchen scraps to enrich the soil with nitrogen and improve drainage for the perennials. ",
    "music": "The orchestra tuned to the oboe before the violins introduced the melody in a minor key. "
             "A fugue weaves several voices together, each entering with the same subject at a different pitch. "
             "The conductor shaped the dynamics of the symphony from a whisper to a thunderous crescendo. ",
    "finance": "Compound interest grows an investment exponentially because each period earns interest on previous interest. "
               "A diversified portfolio spreads risk across stocks, bonds, and index funds to reduce volatility. "
               "The central bank raised interest rates to curb inflation, which increased mortgage payments for borrowers. ",
}


def topic_paragraphs(topic: str, count: int = 12) -> list[str]:
    base = TOPICS[topic]
    return [f"Section {i + 1} on {topic}. " + base * 2 for i in range(count)]


def topic_markdown(topic: str, sections: int = 4, per_section: int = 3) -> str:
    out = [f"# {topic.title()} handbook", ""]
    for s in range(sections):
        out += [f"## Part {s + 1}", ""]
        out += [p + "\n" for p in topic_paragraphs(topic, per_section)]
    return "\n".join(out)
